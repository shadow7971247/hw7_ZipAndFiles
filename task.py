import os
import zipfile
import pytest
from io import BytesIO
import openpyxl
import csv
from pypdf import PdfReader


@pytest.fixture
def zip_archive():
    test_dir = os.path.dirname(__file__)
    tmp_dir = os.path.join(test_dir, "tmp")

    pdf_file = os.path.join(tmp_dir, "document.pdf")
    xlsx_file = os.path.join(tmp_dir, "table.xlsx")
    csv_file = os.path.join(tmp_dir, "data.csv")

    missing = [
        f for f in [pdf_file, xlsx_file, csv_file] if not os.path.exists(f)
    ]
    if missing:
        pytest.skip(f"В папке tmp нет файлов: {missing}")

    archive_path = os.path.join(test_dir, "test_files.zip")
    with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as zipf:
        zipf.write(pdf_file, arcname="document.pdf")
        zipf.write(xlsx_file, arcname="table.xlsx")
        zipf.write(csv_file, arcname="data.csv")

    yield archive_path

    if os.path.exists(archive_path):
        os.remove(archive_path)


def test_pdf_has_pages(zip_archive):
    with zipfile.ZipFile(zip_archive, "r") as zipf:
        pdf_bytes = zipf.read("document.pdf")
        reader = PdfReader(BytesIO(pdf_bytes))
        page_count = len(reader.pages)
        print(f"PDF: {page_count} стр.")
        assert page_count > 0


def test_excel_has_rows(zip_archive):
    with zipfile.ZipFile(zip_archive, "r") as zipf:
        xlsx_bytes = zipf.read("table.xlsx")
        wb = openpyxl.load_workbook(
            BytesIO(xlsx_bytes), read_only=True, data_only=True
        )
        sheet = wb.active
        rows, cols = sheet.max_row, sheet.max_column
        print(f"Excel: {rows} строк, {cols} столбцов")
        assert rows > 0 and cols > 0
        if rows and cols:
            assert sheet.cell(1, 1).value is not None


def test_csv_has_rows(zip_archive):
    with zipfile.ZipFile(zip_archive, "r") as zipf:
        csv_bytes = zipf.read("data.csv")
        encodings = ["utf-8", "utf-8-sig", "cp1251", "windows-1251"]
        rows = None
        for enc in encodings:
            try:
                text = csv_bytes.decode(enc)
                rows = list(csv.reader(text.splitlines()))
                print(f"CSV ({enc}): {len(rows)} строк")
                break
            except UnicodeDecodeError:
                continue
        if rows is None:
            pytest.fail("Не удалось прочитать CSV")
        assert len(rows) > 0
        assert len(rows) >= 2
