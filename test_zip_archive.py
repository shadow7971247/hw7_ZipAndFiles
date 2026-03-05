import zipfile
from io import BytesIO
import csv
import openpyxl
from pypdf import PdfReader


def test_pdf_has_pages_and_text(zip_archive):
    """Проверяет, что PDF содержит хотя бы одну страницу и текст не пуст."""
    with zipfile.ZipFile(zip_archive, "r") as zipf:
        pdf_bytes = zipf.read("document.pdf")
        reader = PdfReader(BytesIO(pdf_bytes))
        assert len(reader.pages) > 0
        text = reader.pages[0].extract_text()
        assert text.strip(), "Текст в PDF отсутствует"


def test_excel_has_rows_and_first_cell(zip_archive):
    """Проверяет, что Excel содержит строки и первая ячейка не пуста."""
    with zipfile.ZipFile(zip_archive, "r") as zipf:
        xlsx_bytes = zipf.read("table.xlsx")
        wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
        sheet = wb.active
        assert sheet.max_row > 0
        assert sheet.max_column > 0
        first_cell = sheet.cell(1, 1).value
        assert first_cell is not None, "Первая ячейка пустая"


def test_csv_has_rows_and_header(zip_archive):
    """Проверяет, что CSV содержит заголовок и хотя бы одну строку данных."""
    with zipfile.ZipFile(zip_archive, "r") as zipf:
        csv_bytes = zipf.read("data.csv")
        encodings = ["utf-8", "utf-8-sig", "cp1251", "windows-1251"]
        rows = None
        for enc in encodings:
            try:
                text = csv_bytes.decode(enc)
                rows = list(csv.reader(text.splitlines()))
                break
            except UnicodeDecodeError:
                continue
        assert rows is not None, "Не удалось прочитать CSV"
        assert len(rows) >= 2, "В CSV должен быть заголовок и хотя бы одна строка данных"
        # Проверяем, что заголовок не пустой
        assert all(rows[0]), "Заголовок содержит пустые поля"
        # Проверяем, что первая строка данных не пустая
        assert any(rows[1]), "Первая строка данных пуста"